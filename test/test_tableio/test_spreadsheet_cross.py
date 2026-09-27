#! /usr/bin/env python3
"""Cross-backend tests for spreadsheet TableIO implementations."""

# Copyright (c) 2026 Tom Björkholm
# MIT License

import re
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pytest import CaptureFixture

from tableio.tableio import Box, FileAccess, TableIO
from tableio.tableio_excel_openpyxl import TableIOExcelOpenPyXL
from tableio.tableio_excel_pylightxl import TableIOExcelPylightxl
from tableio.tableio_excel_xlsxwriter import TableIOExcelXlsxWriter
from tableio.tableio_ods_odfdo import TableIOOdsOdfdo
from tableio.value_type import ListData, Value

from .check_capsys import check_capsys


_READ_WRITE = [TableIOExcelOpenPyXL, TableIOExcelPylightxl, TableIOOdsOdfdo]
_FILTERING = [TableIOExcelOpenPyXL, TableIOExcelXlsxWriter, TableIOOdsOdfdo]
_UPDATE_FILTERING = [TableIOExcelOpenPyXL, TableIOOdsOdfdo]
_EQUALS_DATA: ListData[Value] = [['head', '=x'], ['=1+1', '=A1'],
                                 ['http://a.b', 'plain']]
_DATA: ListData[Value] = [['a', 'b'], [1, 2]]
_WRITE_READ_PAIRS = [
    (TableIOExcelOpenPyXL, TableIOExcelOpenPyXL),
    (TableIOExcelOpenPyXL, TableIOExcelPylightxl),
    (TableIOExcelPylightxl, TableIOExcelPylightxl),
    (TableIOExcelPylightxl, TableIOExcelOpenPyXL),
    (TableIOExcelXlsxWriter, TableIOExcelOpenPyXL),
    (TableIOExcelXlsxWriter, TableIOExcelPylightxl),
    (TableIOOdsOdfdo, TableIOOdsOdfdo)]


def _class_id(tableio_class: type[TableIO]) -> str:
    """Return a short test id for one TableIO class."""
    return tableio_class.__name__.removeprefix('TableIO')


def _excel_filters(file_path: Path) -> dict[str, dict[str, str]]:
    """Return the Excel table names and ranges per worksheet."""
    workbook = load_workbook(file_path)
    ret = {worksheet.title: {name: worksheet.tables[name].ref
                             for name in worksheet.tables}
           for worksheet in workbook.worksheets}
    workbook.close()
    return ret


def _ods_filters(file_path: Path) -> dict[str, dict[str, str]]:
    """Return the ODS database range names and ranges per table."""
    with ZipFile(file_path) as zip_file:
        content = zip_file.read('content.xml').decode('utf-8')
    ret: dict[str, dict[str, str]] = {}
    pattern = r'table:name="([^"]*)"[^>]*target-range-address="([^"]*)"'
    for name, address in re.findall(pattern, content):
        sheet, first, _, last = re.split(r'[.:]', address)
        ret.setdefault(sheet, {})[name] = f'{first}:{last}'
    return ret


def _filters(file_path: Path) -> dict[str, dict[str, str]]:
    """Return the filtered ranges per sheet of one spreadsheet file."""
    if file_path.suffix == '.ods':
        return _ods_filters(file_path)
    return _excel_filters(file_path)


@pytest.mark.parametrize(('writer', 'reader'), _WRITE_READ_PAIRS,
                         ids=_class_id)
def test_equals_text_is_text(writer: type[TableIO], reader: type[TableIO],
                             capsys: CaptureFixture[str]) -> None:
    """Strings starting with '=' are written and read as text."""
    with TemporaryDirectory() as temp_dir:
        with writer(Path(temp_dir) / 'equals', FileAccess.CREATE) as out:
            out.write_table_listdata(_EQUALS_DATA, filtered_data_range=True)
        with reader(out.file_name, FileAccess.READ) as table_io:
            assert table_io.read_table_listdata().data == _EQUALS_DATA
    check_capsys(capsys)


@pytest.mark.parametrize('tableio_class', _READ_WRITE, ids=_class_id)
def test_equals_in_session(tableio_class: type[TableIO],
                           capsys: CaptureFixture[str]) -> None:
    """Strings starting with '=' read back as text in the same session."""
    with TemporaryDirectory() as temp_dir:
        with tableio_class(Path(temp_dir) / 'equals',
                           FileAccess.CREATE) as table_io:
            table_io.write_cells(_EQUALS_DATA, Box(0, 0, 3, 2))
            assert table_io.read_cells(Box(0, 0, 3, 2)) == _EQUALS_DATA
    check_capsys(capsys)


@pytest.mark.parametrize('tableio_class', _FILTERING, ids=_class_id)
def test_filter_names_sheets(tableio_class: type[TableIO],
                             capsys: CaptureFixture[str]) -> None:
    """Filter names are unique in the workbook; overwrites stay per sheet."""
    with TemporaryDirectory() as temp_dir:
        with tableio_class(Path(temp_dir) / 'filters',
                           FileAccess.CREATE) as table_io:
            first_sheet = table_io.current_sheet_name()
            table_io.write_table_listdata(_DATA, filtered_data_range=True)
            table_io.select_sheet('Two', create=True)
            table_io.write_table_listdata(_DATA, filtered_data_range=True)
            table_io.write_table_listdata([['c', 'd'], [3, 4]],
                                          filtered_data_range=True,
                                          box=Box(0, 0, None, None))
        assert _filters(Path(table_io.file_name)) == {
            first_sheet: {'TableIOFilter_1': 'A1:B2'},
            'Two': {'TableIOFilter_2': 'A1:B2'}}
    check_capsys(capsys)


@pytest.mark.parametrize('tableio_class', _UPDATE_FILTERING, ids=_class_id)
def test_filter_names_update(tableio_class: type[TableIO],
                             capsys: CaptureFixture[str]) -> None:
    """UPDATE adds filters with names not used on any existing sheet."""
    with TemporaryDirectory() as temp_dir:
        with tableio_class(Path(temp_dir) / 'filters',
                           FileAccess.CREATE) as table_io:
            first_sheet = table_io.current_sheet_name()
            table_io.write_table_listdata(_DATA, filtered_data_range=True)
        file_path = Path(table_io.file_name)
        with tableio_class(file_path, FileAccess.UPDATE) as table_io:
            table_io.select_sheet('Two', create=True)
            table_io.write_table_listdata(_DATA, filtered_data_range=True)
        assert _filters(file_path) == {
            first_sheet: {'TableIOFilter_1': 'A1:B2'},
            'Two': {'TableIOFilter_2': 'A1:B2'}}
    check_capsys(capsys)


@pytest.mark.parametrize('tableio_class', _READ_WRITE, ids=_class_id)
def test_dict_read_after_end(tableio_class: type[TableIO],
                             capsys: CaptureFixture[str]) -> None:
    """Dict reads after the last table, or of an empty sheet, are empty."""
    with TemporaryDirectory() as temp_dir:
        with tableio_class(Path(temp_dir) / 'end',
                           FileAccess.CREATE) as table_io:
            table_io.write_heading('Only heading')
            table_io.select_sheet('Empty', create=True)
            assert table_io.read_table_dictdata().data == []
            table_io.select_sheet('Two', create=True)
            table_io.write_table_listdata(_DATA)
            assert table_io.read_table_dictdata().data == [{'a': 1, 'b': 2}]
            assert table_io.read_table_dictdata().data == []
    check_capsys(capsys)


@pytest.mark.parametrize('tableio_class', _READ_WRITE, ids=_class_id)
def test_read_moves_write(tableio_class: type[TableIO],
                          capsys: CaptureFixture[str]) -> None:
    """A read moves the default write position to after the rows read."""
    tables: list[ListData[Value]] = [[[name, name], [1, 2]]
                                     for name in ('x', 'y', 'z')]
    new_table: ListData[Value] = [['new', 'new'], [9, 9]]
    with TemporaryDirectory() as temp_dir:
        with tableio_class(Path(temp_dir) / 'cursor',
                           FileAccess.CREATE) as table_io:
            for table in tables:
                table_io.write_table_listdata(table)
            table_io.read_table_listdata(box=Box(0, 0, 2, 2))
            table_io.write_table_listdata(new_table)
        with tableio_class(table_io.file_name, FileAccess.READ) as table_io:
            assert [table_io.read_table_listdata().data
                    for _ in tables] == [tables[0], new_table, tables[2]]
    check_capsys(capsys)

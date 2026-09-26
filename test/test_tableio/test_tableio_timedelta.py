#! /usr/local/bin/python3
"""Tests for writing and reading timedelta values with all backends."""

# Copyright (c) 2026 Tom Björkholm
# MIT License

from datetime import timedelta
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_timedelta_format
from openxml_audit import OpenXmlValidator  # type: ignore[import-untyped]
from tableio import create_tableio, filter_args_tableio, parse_timedelta
from tableio.optional_args import OptionalArgsDict, TimeDeltaFallback
from tableio.tableio import FileAccess, TableIO
from tableio.tableio_csv import TableIOCsv
from tableio.tableio_excel_openpyxl import TableIOExcelOpenPyXL
from tableio.tableio_excel_pylightxl import TableIOExcelPylightxl
from tableio.tableio_excel_xlsxwriter import TableIOExcelXlsxWriter
from tableio.tableio_mformat import TableIOMformatHtml, TableIOMformatMd, \
    TableIOMformatRst, TableIOMformatTxt
from tableio.tableio_mformatbased import TableIOMformatBased
from tableio.tableio_ods_odfdo import TableIOOdsOdfdo, _timedelta_from_odf
from tableio.tableio_types import Box
from tableio.value_type import Fmt, ListData, Value, ValueFmt
from .odf_validation_helper import odf_validation_result

_F = TimeDeltaFallback
_VALUES: list[timedelta] = [
    timedelta(0), timedelta(hours=26, minutes=3, seconds=4),
    timedelta(seconds=1.5), timedelta(hours=-1),
    timedelta(days=15, milliseconds=250), timedelta(weeks=-2, seconds=-7)]
"""Durations with millisecond precision (openpyxl reads milliseconds)."""
_EXCEL_WRITERS: list[type[TableIO]] = [
    TableIOExcelOpenPyXL, TableIOExcelPylightxl, TableIOExcelXlsxWriter]
_PAIRS = [pytest.param(writer, reader,
                       id=f'{writer.__name__}-{reader.__name__}')
          for writer in _EXCEL_WRITERS
          for reader in [TableIOExcelOpenPyXL, TableIOExcelPylightxl]] + \
    [pytest.param(TableIOOdsOdfdo, TableIOOdsOdfdo, id='ods')]


def _table() -> ListData[Value]:
    """Return a table with a header row and one timedelta per row."""
    header: list[Value] = ['name', 'duration']
    return [header] + [[f'value {index}', value]
                       for index, value in enumerate(_VALUES)]


def _written_file(tableio_class: type[TableIO], tmp_path: Path,
                  fallback: Optional[TimeDeltaFallback] = None) -> Path:
    """Write the timedelta table and return the written file path."""
    file_path = Path(tableio_class.file_name_with_extension(
        tmp_path / 'durations', tableio_class.file_name_extension()))
    if fallback is None:
        table_io = tableio_class(file_path, FileAccess.CREATE)
    else:
        assert issubclass(tableio_class, (TableIOCsv, TableIOMformatBased))
        table_io = tableio_class(file_path, FileAccess.CREATE,
                                 timedelta_fallback=fallback)
    with table_io:
        table_io.write_table_listdata(_table())
    return file_path


@pytest.mark.parametrize(('writer', 'reader'), _PAIRS)
def test_native_round_trip(writer: type[TableIO], reader: type[TableIO],
                           tmp_path: Path) -> None:
    """Spreadsheet backends store and read timedelta values natively."""
    file_path = _written_file(writer, tmp_path)
    with reader(file_path, FileAccess.READ) as table_io:
        dict_data = table_io.read_table_dictdata().data
        cells = table_io.read_cells(Box(top=1, left=1, bottom=3, right=2))
        found = table_io.find_value(timedelta(seconds=1.5))
    assert [row['duration'] for row in dict_data] == _VALUES
    assert all(isinstance(row['duration'], timedelta) for row in dict_data)
    assert cells == [[_VALUES[0]], [_VALUES[1]]]
    assert found == Box(top=3, left=1, bottom=4, right=2)


@pytest.mark.parametrize('writer', _EXCEL_WRITERS)
def test_excel_duration_format(writer: type[TableIO], tmp_path: Path) -> None:
    """Excel backends use a duration number format for timedelta."""
    workbook = load_workbook(_written_file(writer, tmp_path))
    worksheet = workbook.active
    assert worksheet is not None
    formats = [worksheet.cell(row=row, column=2).number_format
               for row in range(2, len(_VALUES) + 2)]
    assert all(is_timedelta_format(fmt) for fmt in formats), formats
    assert set(formats) == {'[hh]:mm:ss'}


@pytest.mark.parametrize('writer', _EXCEL_WRITERS + [TableIOOdsOdfdo])
def test_validator_clean(writer: type[TableIO], tmp_path: Path) -> None:
    """Files with timedelta values pass the file format validators."""
    file_path = _written_file(writer, tmp_path)
    if writer is TableIOOdsOdfdo:
        result = odf_validation_result(file_path)
    else:
        result = OpenXmlValidator().validate(file_path)
    assert result.is_valid, [str(error) for error in result.errors]


def test_ods_native_markup(tmp_path: Path) -> None:
    """ODS stores timedelta as time-value durations with fractions."""
    office = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'
    with ZipFile(_written_file(TableIOOdsOdfdo, tmp_path)) as zip_file:
        root = ET.fromstring(zip_file.read('content.xml'))
    durations = [element.get(f'{office}time-value')
                 for element in root.iter()
                 if element.get(f'{office}value-type') == 'time']
    assert durations == ['PT00H00M00S', 'PT26H03M04S', 'PT00H00M01.5S',
                         '-PT01H00M00S', 'PT360H00M00.25S', '-PT336H00M07S']


_ODS_NS = {'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
           'style': 'urn:oasis:names:tc:opendocument:xmlns:style:1.0',
           'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
           'number': 'urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0'}


def _ods_display(file_path: Path) -> list[tuple[str, Optional[ET.Element]]]:
    """Return (value type, time data style) for each ODS data cell."""
    ns = _ODS_NS
    with ZipFile(file_path) as zip_file:
        root = ET.fromstring(zip_file.read('content.xml'))
    styles = {style.get(f'{{{ns["style"]}}}name'): style
              for style in root.iterfind('.//style:style', ns)}
    data_styles = {style.get(f'{{{ns["style"]}}}name'): style
                   for style in root.iterfind('.//number:time-style', ns)}
    ret: list[tuple[str, Optional[ET.Element]]] = []
    for cell in root.iterfind('.//table:table-cell[@office:value-type]', ns):
        style = styles.get(cell.get(f'{{{ns["table"]}}}style-name'))
        data_name = None if style is None else \
            style.get(f'{{{ns["style"]}}}data-style-name')
        ret.append((cell.get(f'{{{ns["office"]}}}value-type', ''),
                    data_styles.get(data_name or '')))
    return ret


def test_ods_duration_display(tmp_path: Path) -> None:
    """ODS duration cells display as '[HH]:MM:SS', also when formatted."""
    file_path = tmp_path / 'display.ods'
    with TableIOOdsOdfdo(file_path, FileAccess.CREATE) as table_io:
        table_io.write_table_listdata(
            [[ValueFmt('text', Fmt()),
              ValueFmt(timedelta(hours=51), Fmt(bold=True))],
             [ValueFmt(timedelta(seconds=-237.5), Fmt()), ValueFmt(7, Fmt())]])
    display = _ods_display(file_path)
    assert [value_type for value_type, _ in display] == \
        ['string', 'time', 'time', 'float']
    assert display[0][1] is None and display[3][1] is None
    number = f'{{{_ODS_NS["number"]}}}'
    for _, data_style in display[1:3]:
        assert data_style is not None
        assert data_style.get(f'{number}truncate-on-overflow') == 'false'
        assert [(part.tag, part.get(f'{number}style'))
                for part in data_style if part.tag != f'{number}text'] == \
            [(f'{number}hours', 'long'), (f'{number}minutes', 'long'),
             (f'{number}seconds', 'long')]


@pytest.mark.parametrize(
    ('text', 'expected'),
    [('PT26H03M04S', timedelta(hours=26, minutes=3, seconds=4)),
     ('PT00H00M01.500S', timedelta(seconds=1.5)),
     ('PT1,25S', timedelta(seconds=1.25)),
     ('-PT01H00M00S', timedelta(hours=-1)),
     ('P2DT3H', timedelta(days=2, hours=3)),
     ('P1D', timedelta(days=1)),
     ('PT', timedelta(0)),
     ('P1Y', None), ('P1W', None), ('1:00:00', None), ('', None),
     (None, None)])
def test_ods_duration_text(text: Optional[str],
                           expected: Optional[timedelta]) -> None:
    """ODF durations, including fractional seconds, are parsed correctly."""
    assert _timedelta_from_odf(text) == expected


@pytest.mark.parametrize(
    ('fallback', 'expected_rows'),
    [(None, ['00:00:00', '26:03:04', '00:00:01.5', '-01:00:00',
             '360:00:00.25', '-336:00:07']),
     (_F.HMS_STRING, ['00:00:00', '26:03:04', '00:00:01.5', '-01:00:00',
                      '360:00:00.25', '-336:00:07']),
     (_F.FLOATSECONDS, ['0.0', '93784.0', '1.5', '-3600.0', '1296000.25',
                        '-1209607.0']),
     (_F.DHMS_STRING, ['00:00:00', '1 d 02:03:04', '00:00:01.5',
                       '-01:00:00', '15 d 00:00:00.25', '-14 d 00:00:07']),
     (_F.WDHMS_STRING_LONG, ['00:00:00', '1 day 02:03:04', '00:00:01.5',
                             '-01:00:00', '2 weeks 1 day 00:00:00.25',
                             '-2 weeks 00:00:07'])])
def test_csv_fallback(fallback: Optional[TimeDeltaFallback],
                      expected_rows: list[str], tmp_path: Path) -> None:
    """CSV writes timedelta as fallback, and parse_timedelta reads it."""
    file_path = _written_file(TableIOCsv, tmp_path, fallback)
    lines = file_path.read_text(encoding='utf-8').splitlines()
    expected_lines = [f'"value {index}","{text}"'
                      for index, text in enumerate(expected_rows)]
    assert lines[1:len(_VALUES) + 1] == expected_lines
    with TableIOCsv(file_path, FileAccess.READ) as table_io:
        data = table_io.read_table_listdata().data
    durations = [row[1] for row in data[1:]]
    assert all(isinstance(duration, str) for duration in durations)
    assert [parse_timedelta(str(duration)) for duration in durations] == \
        _VALUES


@pytest.mark.parametrize('tableio_class', [TableIOMformatMd,
                                           TableIOMformatHtml,
                                           TableIOMformatTxt,
                                           TableIOMformatRst])
@pytest.mark.parametrize(('fallback', 'expected'),
                         [(None, '26:03:04'), (_F.FLOATSECONDS, '93784.0'),
                          (_F.DHMS_STRING_LONG, '1 day 02:03:04')])
def test_mformat_fallback(tableio_class: type[TableIO],
                          fallback: Optional[TimeDeltaFallback], expected: str,
                          tmp_path: Path) -> None:
    """Mformat based writers write timedelta as fallback text."""
    file_path = _written_file(tableio_class, tmp_path, fallback)
    text = file_path.read_text(encoding='utf-8')
    assert expected in text
    assert '1 day, 2:03:04' not in text


def test_factory_passes_arg(tmp_path: Path) -> None:
    """The factory passes timedelta_fallback to backends that accept it."""
    args: OptionalArgsDict = {'timedelta_fallback': _F.WDHMS_STRING}
    table_io = create_tableio('CSV', tmp_path / 'factory', FileAccess.CREATE,
                              args=args)
    assert isinstance(table_io, TableIOCsv)
    assert table_io.timedelta_fallback == _F.WDHMS_STRING
    assert filter_args_tableio(args, 'md', None) == args
    assert filter_args_tableio(args, 'Excel', None) == {}
    assert filter_args_tableio(args, 'ODS', None) == {}
